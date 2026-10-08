#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_plz.py – liest den PLZ-Checker (Excel) und schreibt die Daten in index.html.

Aufruf:
    python3 build_plz.py PLZ-Checker.xlsx            # schreibt index.html neu
    python3 build_plz.py PLZ-Checker.xlsx --check    # nur prüfen, nichts schreiben

Erwartete Excel-Struktur (wie "PLZ-Checker_General_HVAC"):
    Tabelle2: PLZ von | PLZ bis | Verkäufer | Planerberater | RC | ID kfm | Techn. Innendienst
    Tabelle3: Name | Tel | E-Mail | RC
    Datum:    Datenstand als Text (z. B. "02.03.26")

Der Datenblock in index.html liegt zwischen /*DATA-START*/ und /*DATA-END*/.
"""
import json
import re
import sys
from datetime import datetime

from openpyxl import load_workbook

# Hier können bekannte Tippfehler aus der Excel korrigiert werden, ohne die
# Excel anzufassen – z. B. E-Mail-Adressen mit Umlauten. Leer = keine Korrektur.
FIXES = {
    # "jürgen.kunzi@de.generalww.com": "juergen.kunzi@de.generalww.com",
}


def clean_phone(raw):
    """Gibt (anzeige, tel-link) zurück oder (None, None)."""
    if raw is None:
        return None, None
    s = str(raw).strip()
    digits = re.sub(r"\D", "", s)
    if not digits:
        return None, None
    # Anzeige: Vorwahl und Rufnummer mit einem Leerzeichen trennen.
    m = re.match(r"^\s*(\d+)\s*/\s*(\d+)\s*$", s)
    if m:
        display = m.group(1) + " " + m.group(2)
    else:
        display = digits
        # unformatierte Nummern: bekannte Vorwahlen aus derselben Liste erkennen
        for code in ("06102", "0711", "0211", "0176", "089", "030"):
            if digits.startswith(code):
                display = code + " " + digits[len(code):]
                break
    return display, "+49" + digits[1:] if digits.startswith("0") else digits


def read_excel(path):
    wb = load_workbook(path, read_only=True, data_only=True)

    ranges = []
    for row in wb["Tabelle2"].iter_rows(min_row=2, values_only=True):
        if row[0] is None or not isinstance(row[0], (int, float)):
            continue
        ranges.append({
            "from": int(row[0]),
            "to": int(row[1]),
            "sales": (row[2] or "").strip(),
            "rc": (row[4] or "").strip(),
            "kfm": (row[5] or "").strip(),
            "tech": (row[6] or "").strip(),
        })
    ranges.sort(key=lambda r: r["from"])

    contacts = {}
    for row in wb["Tabelle3"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        name = str(row[0]).strip()
        display, tel = clean_phone(row[1])
        mail = str(row[2]).strip() if row[2] else None
        if mail:
            mail = FIXES.get(mail, mail)
        contacts[name] = {
            "tel": tel,
            "telDisplay": display,
            "mail": mail,
            "rc": str(row[3]).strip() if row[3] else None,
        }

    stand = None
    if "Datum" in wb.sheetnames:
        for row in wb["Datum"].iter_rows(values_only=True):
            if row and row[0]:
                stand = str(row[0]).strip()
                break
    return ranges, contacts, stand


def check(ranges, contacts):
    problems = []
    prev = None
    for r in ranges:
        if r["to"] < r["from"]:
            problems.append("Bereich falsch herum: %d–%d" % (r["from"], r["to"]))
        if prev is not None:
            if r["from"] <= prev["to"]:
                problems.append("Überlappung: %d–%d und %d–%d" % (prev["from"], prev["to"], r["from"], r["to"]))
            elif r["from"] != prev["to"] + 1:
                problems.append("Lücke: %d–%d ist keinem Verkäufer zugeordnet" % (prev["to"] + 1, r["from"] - 1))
        prev = r
    names = set()
    for r in ranges:
        for key in ("sales", "kfm", "tech"):
            if r[key]:
                names.add(r[key])
    for n in sorted(names):
        if n not in contacts:
            problems.append("Keine Kontaktdaten in Tabelle3 für: %s" % n)
    for n, c in contacts.items():
        if c["mail"] and not re.match(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$", c["mail"]):
            problems.append("E-Mail enthält Sonderzeichen (bitte prüfen): %s -> %s" % (n, c["mail"]))
        if not c["tel"]:
            problems.append("Keine Telefonnummer: %s" % n)
    return problems


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    xlsx = sys.argv[1]
    only_check = "--check" in sys.argv

    ranges, contacts, stand = read_excel(xlsx)
    problems = check(ranges, contacts)

    print("Bereiche: %d  (von %05d bis %05d)" % (len(ranges), ranges[0]["from"], ranges[-1]["to"]))
    print("Kontakte: %d" % len(contacts))
    print("Datenstand laut Excel: %s" % stand)
    if problems:
        print("\nHinweise aus der Prüfung:")
        for p in problems:
            print("  - " + p)
    if only_check:
        return

    data = {
        "stand": stand,
        "built": datetime.now().strftime("%Y-%m-%d"),
        "ranges": ranges,
        "contacts": contacts,
    }
    block = "/*DATA-START*/var DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";/*DATA-END*/"

    with open("index.html", "r", encoding="utf-8") as f:
        html = f.read()
    new_html, n = re.subn(r"/\*DATA-START\*/.*?/\*DATA-END\*/", lambda m: block, html, flags=re.S)
    if n != 1:
        print("FEHLER: Datenblock /*DATA-START*/ … /*DATA-END*/ nicht (genau einmal) in index.html gefunden.")
        sys.exit(2)
    with open("index.html", "w", encoding="utf-8") as f:
        f.write(new_html)
    print("\nindex.html aktualisiert.")


if __name__ == "__main__":
    main()
